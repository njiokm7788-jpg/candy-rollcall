"""导入导出：CSV / XLSX 名单（§8.2）。

- CSV 优先 utf-8-sig，提供 GB18030 编码选项；学号按字符串处理。
- XLSX 用 openpyxl；学号列读取时若为数值不凭空猜前导零。
- 导入返回预览（有效条数、空姓名、疑似重复、超额），不写库。
"""
from __future__ import annotations

import csv
import io
import re
from typing import Dict, List, Optional, Tuple

from ..core.models import Student, ValidationError, MAX_STUDENTS

SUPPORTED_CSV_ENCODINGS = ("utf-8-sig", "gb18030", "utf-8")


# ---------------- 读取原始行 ----------------

def read_csv_rows(content: bytes, encoding: str = "utf-8-sig") -> List[List[str]]:
    """解码并解析 CSV，返回每行单元格（字符串）。"""
    if encoding not in SUPPORTED_CSV_ENCODINGS:
        raise ValidationError(f"不支持的编码: {encoding}")
    try:
        text = content.decode(encoding)
    except (UnicodeDecodeError, LookupError) as e:
        raise ValidationError(f"无法以 {encoding} 解码文件: {e}")
    reader = csv.reader(io.StringIO(text))
    return [row for row in reader]


def read_xlsx_rows(path: str) -> List[List[str]]:
    """openpyxl 读取第一个工作表，所有单元格转为字符串（保留 "08" 等文本格式）。"""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        rows: List[List[str]] = []
        for row in ws.iter_rows(values_only=True):
            cells: List[str] = []
            for cell in row:
                if cell is None:
                    cells.append("")
                elif isinstance(cell, float) and cell.is_integer():
                    # Excel 数值 8 -> "8"；前导零本已丢失，不凭空猜测（§8.2）
                    cells.append(str(int(cell)))
                else:
                    cells.append(str(cell))
            # 去除行尾全空
            while cells and cells[-1] == "":
                cells.pop()
            rows.append(cells)
        return rows
    finally:
        wb.close()


# ---------------- 列映射推导（§8.2） ----------------

# 学号列的表头关键词；“序号”必须排除，否则本工具自己导出的表格会错位（AT-14）
_NAME_HEADER_PATTERNS = (r"姓名", r"名字", r"学生", r"name")
_NO_HEADER_PATTERNS = (r"学号", r"学籍号", r"编号", r"^no\.?$", r"^id$")


def _pick_column(header: List[str], patterns: Tuple[str, ...]) -> int:
    """按表头文本定位列，未命中返回 -1。"""
    for i, cell in enumerate(header):
        for pattern in patterns:
            if re.search(pattern, cell, re.IGNORECASE):
                return i
    return -1


def derive_mapping(rows: List[List[str]]) -> Tuple[Dict[str, int], bool]:
    """推导导入列映射（§8.2）。

    有表头时按表头文本定位列——这样本工具自己导出的“序号/姓名/学号”三列
    表格可以直接重新导入（AT-14）；无表头时回退第 1 列姓名、第 2 列学号。
    返回 (mapping, has_header)；无学号列时用 -1 表示“不存在该列”。
    """
    if not rows:
        return {"name": 0, "studentNo": -1}, False
    header = [str(cell) for cell in rows[0]]
    has_header = any(re.search(r"姓名|学号|name", cell, re.IGNORECASE) for cell in header)
    if not has_header:
        return {"name": 0, "studentNo": 1}, False
    name_col = _pick_column(header, _NAME_HEADER_PATTERNS)
    if name_col < 0:
        name_col = 0
    return {"name": name_col, "studentNo": _pick_column(header, _NO_HEADER_PATTERNS)}, True


# ---------------- 行 → 学生草稿 ----------------

def rows_to_drafts(rows: List[List[str]], mapping: Dict[str, int],
                   has_header: bool = False) -> List[dict]:
    """按列映射把原始行转为草稿。mapping: {"name": col, "studentNo": col 或 -1}。"""
    name_col = mapping.get("name", 0)
    no_col = mapping.get("studentNo", -1)
    drafts: List[dict] = []
    data_rows = rows[1:] if has_header else rows
    for row in data_rows:
        if not any(cell.strip() for cell in row):
            continue  # 跳过空行
        name = row[name_col].strip() if name_col < len(row) else ""
        no = ""
        if no_col is not None and 0 <= no_col < len(row):
            no = row[no_col].strip()
        drafts.append({"name": name, "studentNo": no or None})
    return drafts


def parse_paste_text(text: str) -> List[dict]:
    """批量粘贴：每行一个姓名或 “学号+姓名”/“姓名+学号” 两列（制表符/逗号/空格分隔）。"""
    drafts: List[dict] = []
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        for sep in ("\t", ","):
            if sep in line:
                parts = [p.strip() for p in line.split(sep) if p.strip()]
                break
        else:
            parts = line.split()
        if len(parts) == 1:
            drafts.append({"name": parts[0], "studentNo": None})
        elif len(parts) >= 2:
            a, b = parts[0], parts[1]
            # 纯数字在前视为学号列，否则第一个视为姓名
            if a.isdigit() and not b.isdigit():
                drafts.append({"name": b, "studentNo": a})
            else:
                drafts.append({"name": a, "studentNo": b if b.isdigit() else None})
    return drafts


# ---------------- 预览与重复检查（§8.3） ----------------

def preview_import(drafts: List[dict], existing: List[Student]) -> dict:
    """返回预览：有效条数、空姓名、超额、疑似重复清单。不写库。"""
    valid: List[dict] = []
    empty_names: List[int] = []
    for i, d in enumerate(drafts):
        if not (d.get("name") or "").strip():
            empty_names.append(i)
        else:
            valid.append(d)

    issues: List[dict] = []
    # 新名单内部重复：同一 (姓名, 学号) 对出现多次 → 报告第二次及以后（高度疑似）
    internal_dup_reported = set()
    seen: Dict[Tuple[str, str], List[int]] = {}
    for i, d in enumerate(valid):
        key = (d["name"], d.get("studentNo") or "")
        if key in seen:
            internal_dup_reported.add(i)
        else:
            seen[key] = [i]

    existing_pairs = {(s.name, s.studentNo or "") for s in existing}
    existing_names = {s.name for s in existing}
    existing_nos = {s.studentNo for s in existing if s.studentNo}

    for i, d in enumerate(valid):
        name = d["name"]
        no = d.get("studentNo") or ""
        key = (name, no)
        kind: Optional[str] = None
        detail = ""
        if i in internal_dup_reported:
            kind = "duplicate_exact"
            detail = "新名单内部重复（同名同号）"
        elif name in existing_names:
            # 同名优先判定：双方无学号时须“人工确认”，不得算作高度疑似（§8.3）
            if no and key in existing_pairs:
                kind = "duplicate_exact"
                detail = "与原名单同名同号"
            elif no:
                kind = "same_name_diff_no"
                detail = "与原名单同名不同学号（可能是同名学生）"
            else:
                kind = "same_name_no_no"
                detail = "与原名单同名且双方都无学号，需人工确认"
        elif no and no in existing_nos:
            kind = "no_conflict"
            detail = "学号与原名单他人冲突（同学号不同姓名）"
        if kind:
            issues.append({"key": f"new:{i}", "index": i, "name": name,
                           "studentNo": no or None, "kind": kind, "detail": detail})

    over_limit = len(valid) > MAX_STUDENTS
    return {
        "totalCount": len(drafts),
        "validCount": len(valid),
        "emptyNameRows": empty_names,
        "issues": issues,
        "overLimit": over_limit,
        "maxStudents": MAX_STUDENTS,
        "drafts": valid,
    }


# ---------------- 导出 ----------------

def export_csv(students: List[Student]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["序号", "姓名", "学号"])
    for i, s in enumerate(students, 1):
        writer.writerow([i, s.name, s.studentNo or ""])
    return buf.getvalue().encode("utf-8-sig")


def export_xlsx(students: List[Student]) -> bytes:
    import io as _io
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "学生名单"
    ws.append(["序号", "姓名", "学号"])
    # 学号列设为文本格式，保住前导零（AT-14）
    for i, s in enumerate(students, 1):
        row_idx = i + 1
        ws.cell(row=row_idx, column=1, value=i)
        ws.cell(row=row_idx, column=2, value=s.name)
        c = ws.cell(row=row_idx, column=3, value=s.studentNo or "")
        c.number_format = "@"
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 12
    buf = _io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
